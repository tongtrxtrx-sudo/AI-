#!/usr/bin/env python3
"""Render the reviewed EPM documentation into editable Chinese Word files.

No network requests or runtime account data are used by this build script.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
MD = MarkdownIt("commonmark").enable("table")
BLUE = "173F6E"
SPECS = [
    ("01-系统研发与交互说明书.md", "01_系统研发与交互说明书", "系统研发与交互说明书", "TY-EPM-SYS-001", "研发架构 · 交互协议 · 数据持久化 · 构建交付", "研发、测试、实施及技术负责人"),
    ("02-产品说明书_含演示案例.md", "02_产品说明书_含演示案例", "产品说明书（含演示案例）", "TY-EPM-PROD-001", "产品模块 · 业务联动 · 演示案例 · 预期结果", "产品负责人、财务负责人及演示评审人员"),
    ("03-用户操作手册.md", "03_用户操作手册", "用户操作手册", "TY-EPM-USER-001", "安装启动 · 字段填写 · 日常操作 · 常见问题", "集团财务、法人会计、预算及经营分析人员"),
]


def font(run, size=None, *, code=False):
    run.font.name = "Consolas" if code else "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    if size:
        run.font.size = Pt(size)


def shade(element, fill):
    s = OxmlElement("w:shd")
    s.set(qn("w:fill"), fill)
    element.append(s)


def field(paragraph, instruction):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), instruction)
    paragraph._p.append(f)


def hyperlink(paragraph, text, url=None, anchor=None, size=None):
    link = OxmlElement("w:hyperlink")
    if anchor:
        link.set(qn("w:anchor"), anchor)
    else:
        rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "245CA8"); props.append(color)
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Arial"); fonts.set(qn("w:eastAsia"), "宋体"); props.append(fonts)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); props.append(sz)
    run.append(props)
    node = OxmlElement("w:t"); node.text = text; node.set(qn("xml:space"), "preserve"); run.append(node)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, tokens, size=None):
    bold = italic = 0
    link_url = None
    link_parts = []
    for token in tokens or []:
        if token.type == "strong_open": bold += 1
        elif token.type == "strong_close": bold -= 1
        elif token.type == "em_open": italic += 1
        elif token.type == "em_close": italic -= 1
        elif token.type == "link_open": link_url, link_parts = token.attrGet("href"), []
        elif token.type == "link_close":
            if link_url:
                if link_url.startswith("#"):
                    hyperlink(paragraph, "".join(link_parts), anchor=link_url[1:], size=size)
                else:
                    hyperlink(paragraph, "".join(link_parts), url=link_url, size=size)
            link_url = None
        elif token.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif token.type in ("text", "code_inline", "html_inline"):
            if link_url:
                link_parts.append(token.content)
                continue
            if token.type == "html_inline" and token.content in ("<br>", "<br/>", "<br />"):
                paragraph.add_run().add_break(); continue
            run = paragraph.add_run(token.content)
            font(run, size or (9 if token.type == "code_inline" else None), code=token.type == "code_inline")
            run.bold = bold > 0
            run.italic = italic > 0


def add_table(doc, rows):
    if not rows:
        return
    columns = max(map(len, rows))
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    samples = [1] * columns
    for row in rows[:50]:
        for i, cell in enumerate(row):
            text = "".join(t.content for t in cell if t.type in ("text", "code_inline"))
            samples[i] = max(samples[i], min(65, len(text)))
    weights = [max(10, s ** .62 * 3.2) for s in samples]
    widths = [17 * w / sum(weights) for w in weights]
    for i, col in enumerate(table.columns): col.width = Cm(widths[i])
    size = 8.3 if columns >= 6 else 9 if columns >= 4 else 9.5
    for ri, row in enumerate(rows):
        cells = table.rows[0].cells if ri == 0 else table.add_row().cells
        trpr = table.rows[ri]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit"); trpr.append(cant_split)
        for i, cell in enumerate(cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                e = OxmlElement("w:" + side); e.set(qn("w:w"), "65"); e.set(qn("w:type"), "dxa"); margins.append(e)
            props.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            inline(p, row[i] if i < len(row) else [], size=size)
            if ri == 0:
                shade(props, "E7EFFA")
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.bold = True; run.font.color.rgb = RGBColor.from_string(BLUE)
            elif ri % 2 == 0:
                shade(props, "F8FAFD")
        if ri == 0:
            repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); trpr.append(repeat)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)


def parse_table(tokens, at):
    rows, current = [], []
    while at < len(tokens):
        t = tokens[at]
        if t.type == "tr_open": current = []
        elif t.type == "inline": current.append(t.children or [])
        elif t.type == "tr_close": rows.append(current)
        elif t.type == "table_close": return rows, at
        at += 1
    return rows, at


def new_document(title, number, subtitle, audience):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.95), Cm(1.85)
    section.left_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(.8)
    section.different_first_page_header_footer = True
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        s = doc.styles[name]
        s.font.name = "Arial"
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        s = doc.styles[name]
        s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(BLUE)
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    r = header.add_run("杭州天元宠物用品有限公司  |  " + title)
    font(r, 8); r.font.color.rgb = RGBColor.from_string("66758D")
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run(number + "  ·  V1.0  ·  第 "), 8)
    field(footer, "PAGE"); font(footer.add_run(" 页 / 共 "), 8); field(footer, "NUMPAGES"); font(footer.add_run(" 页"), 8)
    doc.core_properties.title = "天元EPM · " + title
    doc.core_properties.subject = "杭州天元宠物用品有限公司：中文本机EPM Demo文档"
    doc.core_properties.author = "自研EPM项目文档"
    doc.core_properties.keywords = "集团财务合并,全面预算,经营分析,中文Demo,研发交互,用户手册"
    p = doc.add_paragraph(); p.paragraph_format.space_before = Cm(1.8)
    r = p.add_run("天元 EPM  /  中文本机演示版"); font(r, 12); r.font.color.rgb = RGBColor.from_string("4B6995")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(14)
    r = p.add_run(title); font(r, 25); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph(); font(p.add_run("杭州天元宠物用品有限公司"), 15)
    p = doc.add_paragraph(); r = p.add_run(subtitle); font(r, 11); r.font.color.rgb = RGBColor.from_string("66758D")
    doc.add_paragraph()
    controls = [
        ("文档编号", number), ("文档版本 / 日期", "V1.0 / 2026年10月20日"),
        ("适配软件版本", "20261020-cn-demo1 · 已交付全中文持久化 Demo"),
        ("适用读者", audience),
        ("数据说明", "主公司名称由用户提供；演示子法人、人员、交易与金额均为虚构。"),
    ]
    rows = [[type("T", (), {"type": "text", "content": value})() ] for value in ("文档控制项", "说明")]
    table_rows = [rows] + [[[type("T", (), {"type": "text", "content": value})()] for value in row] for row in controls]
    add_table(doc, table_rows)
    p = doc.add_paragraph("本套说明依据实际交付源码、接口与验收结果整理，明确区分可操作功能、演示计算口径及后续研发建议。")
    p.paragraph_format.space_before = Pt(10)
    shade(p._p.get_or_add_pPr(), "F0F5FC")
    doc.add_page_break()
    return doc


def render(spec):
    filename, stem, title, number, subtitle, audience = spec
    source = ROOT / "content" / filename
    text = source.read_text(encoding="utf-8")
    tokens = MD.parse(text)
    doc = new_document(title, number, subtitle, audience)
    headings = []
    for i, token in enumerate(tokens[:-1]):
        if token.type == "heading_open" and token.tag != "h1":
            headings.append((i, int(token.tag[1:]), tokens[i+1].content, "epm_" + str(len(headings)+1)))
    lookup = {i: (level, text, anchor, n+1) for n, (i, level, text, anchor) in enumerate(headings)}
    doc.add_heading("阅读导航", level=1)
    p = doc.add_paragraph("点击下列章节可跳转。Word/WPS 的标题样式也可用于导航及生成自动目录。")
    for _, level, label, anchor in headings:
        if level > 2: continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        hyperlink(p, label, anchor=anchor, size=10.5)
    doc.add_page_break()
    i, figure = 0, 0
    lists = []
    pending_item = False
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            if t.tag == "h1": i += 3; continue
            level, label, anchor, bookmark_id = lookup[i]
            p = doc.add_heading(level=min(level-1, 4))
            start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), anchor)
            end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_id))
            p._p.append(start); inline(p, tokens[i+1].children); p._p.append(end)
            i += 3; continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            lists.append({"ordered": t.type == "ordered_list_open", "next": int(t.attrGet("start") or 1)})
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif t.type == "list_item_open": pending_item = True
        elif t.type == "paragraph_open":
            children = tokens[i+1].children or []
            images = [x for x in children if x.type == "image"]
            if images:
                for image in images:
                    path = (source.parent / image.attrGet("src")).resolve()
                    if not path.is_file(): raise FileNotFoundError(path)
                    with Image.open(path) as pic: width, height = pic.size
                    display_width = min(16.8, 18.6 * width / height)
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(str(path), width=Cm(display_width))
                    figure += 1
                    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run(f"图 {figure}  {image.content}"); font(r, 9); r.font.color.rgb = RGBColor.from_string("66758D")
            else:
                p = doc.add_paragraph()
                if lists:
                    p.paragraph_format.left_indent = Cm(.5 * len(lists))
                    if pending_item:
                        current = lists[-1]
                        prefix = str(current["next"]) + ". " if current["ordered"] else "• "
                        current["next"] += 1
                        p.paragraph_format.first_line_indent = Cm(-.35)
                        font(p.add_run(prefix))
                inline(p, children)
            pending_item = False
            i += 3; continue
        elif t.type == "table_open":
            rows, i = parse_table(tokens, i)
            add_table(doc, rows)
        elif t.type in ("fence", "code_block"):
            for line in t.content.rstrip("\n").splitlines():
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(.2)
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                shade(p._p.get_or_add_pPr(), "F4F6FA")
                font(p.add_run(line), 8.2, code=True)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif t.type == "hr":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        i += 1
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / (stem + ".docx")
    doc.save(target)
    return {"file": target.name, "source": filename, "sourceSha256": hashlib.sha256(text.encode()).hexdigest(),
            "sha256": hashlib.sha256(target.read_bytes()).hexdigest(), "bytes": target.stat().st_size,
            "headings": len(headings), "tables": len(doc.tables), "figures": figure,
            "chineseCharacters": len(re.findall(r"[\u4e00-\u9fff]", text))}


if __name__ == "__main__":
    records = [render(spec) for spec in SPECS]
    (ROOT / "verification").mkdir(parents=True, exist_ok=True)
    (ROOT / "verification" / "word-build.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(records, ensure_ascii=False))
